# -*- coding: utf-8 -*-
"""
app.services.track_service
==========================
选刊追踪台账服务：管理 CSV/Excel 追踪表同步写入、自动首行冻结与多维条件筛选。
"""
import csv
import os
from app.config import get_track_csv_path, get_track_xlsx_path
from app.core.constants import (
    LEVEL_LABEL,
    LEVEL_ORDER,
    OWN_JOURNALS,
    TRACK_HEADERS,
)
from app.core.rules import index_summary


def build_track_row(journal_rec, kb_rec, catalog, best_level, kb_only=False):
    """由查询结果生成追踪表一行（严格对齐 12 列表头）。"""
    kb = kb_rec or {}
    basic = kb.get("basic") or {}
    if_ = kb.get("impact_factor") or {}
    editor = kb.get("editor_in_chief") or {}
    idx = kb.get("indexes") or {}
    verified = bool(idx.get("indexes_verified"))

    name = journal_rec["name"]
    if best_level:
        level = LEVEL_LABEL.get(best_level, best_level)
    else:
        level = "不符合/待核验"

    # B列: ISSN / CN
    issn_cn = []
    issn_val = journal_rec.get("issn") or basic.get("issn")
    cn_val = journal_rec.get("cn") or basic.get("cn")
    if issn_val:
        issn_cn.append(f"ISSN {issn_val}")
    if cn_val:
        issn_cn.append(f"CN {cn_val}")
    issn_cn_str = " | ".join(issn_cn) if issn_cn else ""

    # J列: 出版方式与版面费
    pub_mode = kb.get("publication_mode") or ""
    apc = kb.get("apc_fee") or ""
    if pub_mode and apc:
        pub_apc_str = f"{pub_mode}；版面费: {apc}"
    elif pub_mode:
        pub_apc_str = pub_mode
    elif apc:
        pub_apc_str = f"版面费: {apc}"
    else:
        pub_apc_str = "待核验" if not verified else "未获取"

    # K列: 投稿方式
    sub_url = kb.get("submission_url") or ("待核验" if not verified else "")

    # L列: 核验日期与来源
    sources = kb.get("sources") or []
    updated_str = kb.get("updated") or ""
    if updated_str and sources:
        verify_str = f"{updated_str}（{len(sources)}个来源）"
    elif updated_str:
        verify_str = updated_str
    else:
        verify_str = "待核验" if not verified else ""

    note = []
    # 风险排查：中科院预警期刊
    from app.services.risk_service import check_cas_warning
    warn = check_cas_warning(name, issn_val)
    if warn:
        note.append(f"⚠️中科院预警期刊({warn['reason']})")

    if name in OWN_JOURNALS:
        note.append("本单位期刊限额1篇")
    if best_level and not kb_only:
        cl = journal_rec.get("level")
        if cl and LEVEL_ORDER.get(best_level, 9) < LEVEL_ORDER.get(cl, 9):
            note.append("按备注1升级认定")
    if kb_only and best_level:
        note.append("目录外按索引规则认定")

    if note and verify_str:
        verify_str = f"{verify_str} [{'; '.join(note)}]"
    elif note:
        verify_str = f"[{'; '.join(note)}]"

    return {
        "期刊名称": name,
        "ISSN / CN": issn_cn_str,
        "收录索引": index_summary(kb_rec),
        "广中医认定级别": level,
        "影响因子(JIF)": (if_.get("jcr_if") or "") if verified and if_ else "待核验",
        "中科院分区": (if_.get("cas_zone") or "") if verified and if_ else "待核验",
        "主编信息": (editor.get("name") or "") if verified and editor else "待核验",
        "收稿方向": (kb.get("scope") or "") if verified and kb.get("scope") else "待核验",
        "审稿周期": (kb.get("review_cycle") or "") if verified and kb.get("review_cycle") else "待核验",
        "出版方式与版面费": pub_apc_str,
        "投稿方式": sub_url,
        "核验日期与来源": verify_str,
    }


def load_track(csv_path=None, user_id=None):
    """加载追踪表记录（支持多用户独立隔离）。"""
    target = csv_path or get_track_csv_path(user_id=user_id)
    if not os.path.exists(target):
        # 若指定了非默认用户但其目录尚无台账，返回空列表
        if user_id:
            return []
        # 兼容性检查：若 root 目录下有旧的 journal_track.csv，读取它
        root_old = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(target))), "journal_track.csv")
        if os.path.exists(root_old):
            target = root_old
        else:
            return []
    with open(target, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    clean_rows = []
    for r in rows:
        row = {}
        for h in TRACK_HEADERS:
            row[h] = r.get(h, "")
        # 兼容旧版本追踪表字段映射
        if not row["主编信息"] and r.get("主编"):
            row["主编信息"] = r.get("主编", "")
        if not row["核验日期与来源"]:
            dt = r.get("核验日期", "")
            cnt = r.get("来源链接数", "")
            rmk = r.get("备注", "")
            parts = []
            if dt and cnt:
                parts.append(f"{dt}（{cnt}个来源）")
            elif dt:
                parts.append(dt)
            if rmk:
                parts.append(f"[{rmk}]")
            row["核验日期与来源"] = " ".join(parts)
        clean_rows.append(row)
    return clean_rows


def save_track(rows, csv_path=None, xlsx_path=None, user_id=None):
    """写 CSV（utf-8-sig）并同步生成带样式的 xlsx（首行冻结与筛选，支持多用户独立隔离）。"""
    c_path = csv_path or get_track_csv_path(user_id=user_id)
    x_path = xlsx_path or get_track_xlsx_path(user_id=user_id)
    os.makedirs(os.path.dirname(c_path), exist_ok=True)
    os.makedirs(os.path.dirname(x_path), exist_ok=True)

    with open(c_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=TRACK_HEADERS)
        w.writeheader()
        w.writerows(rows)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "选刊追踪"
        ws.append(TRACK_HEADERS)
        for r in rows:
            ws.append([r.get(h, "") for h in TRACK_HEADERS])

        # 表头样式 + 冻结 + 筛选
        header_fill = PatternFill("solid", fgColor="2F5597")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(TRACK_HEADERS))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"
        widths = {
            "A": 36,  # 期刊名称
            "B": 24,  # ISSN / CN
            "C": 42,  # 收录索引
            "D": 16,  # 广中医认定级别
            "E": 18,  # 影响因子(JIF)
            "F": 22,  # 中科院分区
            "G": 28,  # 主编信息
            "H": 50,  # 收稿方向
            "I": 42,  # 审稿周期
            "J": 30,  # 出版方式与版面费
            "K": 36,  # 投稿方式
            "L": 22,  # 核验日期与来源
        }
        for col, wd in widths.items():
            ws.column_dimensions[col].width = wd
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        wb.save(x_path)
    except ImportError:
        pass


def upsert_track(row, csv_path=None, xlsx_path=None, user_id=None):
    """幂等写入/更新追踪表。"""
    rows = load_track(csv_path=csv_path, user_id=user_id)
    for r in rows:
        if r.get("期刊名称") == row["期刊名称"]:
            r.update(row)
            break
    else:
        rows.append(row)
    save_track(rows, csv_path=csv_path, xlsx_path=xlsx_path, user_id=user_id)


def delete_track_row(journal_name: str, csv_path=None, xlsx_path=None, user_id=None):
    """从追踪台账中删除指定期刊（支持精确及去除空格比对）。
    返回: (found: bool, remaining_count: int)
    """
    rows = load_track(csv_path=csv_path, user_id=user_id)
    target_name = (journal_name or "").strip().lower()
    new_rows = [r for r in rows if r.get("期刊名称", "").strip().lower() != target_name]
    if len(new_rows) < len(rows):
        save_track(new_rows, csv_path=csv_path, xlsx_path=xlsx_path, user_id=user_id)
        return True, len(new_rows)
    return False, len(rows)


def clear_track(csv_path=None, xlsx_path=None, user_id=None):
    """清空指定用户的全部追踪台账。"""
    save_track([], csv_path=csv_path, xlsx_path=xlsx_path, user_id=user_id)
    return True


def filter_track_rows(track_rows, filter_expr):
    """根据多维复合条件筛选已沉淀的期刊数据。
    支持条件：level(级别), oa(是否OA), scope(收稿方向), apc(版面费), zone(中科院分区)
    """
    if not filter_expr or not str(filter_expr).strip():
        return track_rows

    conds = {}
    for part in filter_expr.split(","):
        part = part.strip()
        if not part:
            continue
        if "=" in part:
            k, v = part.split("=", 1)
            conds[k.strip().lower()] = v.strip()
        elif ":" in part:
            k, v = part.split(":", 1)
            conds[k.strip().lower()] = v.strip()
        else:
            conds[part.strip().lower()] = "true"

    def match_row(r, c_dict):
        for k, v in c_dict.items():
            k = k.lower().strip()
            v = v.strip()
            if k in ("level", "级别"):
                level_str = r.get("广中医认定级别", "")
                if v.upper() in ("A1", "A2", "A3", "A4", "B"):
                    target = f"A类{v[1]}级" if v.upper().startswith("A") else "B类"
                    if target not in level_str:
                        return False
                elif v not in level_str:
                    return False
            elif k in ("oa", "开放获取"):
                pub = r.get("出版方式与版面费", "")
                if v in ("否", "no", "false", "0", "非oa", "非"):
                    if "非oa" not in pub.lower() and "订阅" not in pub and "免版面费" not in pub:
                        return False
                elif v in ("是", "yes", "true", "1", "oa"):
                    if "oa" not in pub.lower() or "非oa" in pub.lower():
                        return False
            elif k in ("scope", "方向", "收稿"):
                if v.lower() not in r.get("收稿方向", "").lower():
                    return False
            elif k in ("zone", "分区"):
                if v not in r.get("中科院分区", ""):
                    return False
            elif k in ("apc", "版面费"):
                pub = r.get("出版方式与版面费", "")
                if v in ("0", "免", "free", "无"):
                    if "免" not in pub and "无版面费" not in pub:
                        return False
                elif v.lower() not in pub.lower():
                    return False
        return True

    return [r for r in track_rows if match_row(r, conds)]


def get_filterable_rows(scope: str = "official", user_id=None):
    """获取供智能筛选的期刊列表。
    scope='official': 广中医官方308号目录136本 + 中文核心/CSCD精选库 (默认)
    scope='track': 仅个人已保存的选刊台账
    scope='all': 全量合并
    """
    track_rows = load_track(user_id=user_id)
    if scope == "track":
        return track_rows

    from app.services.catalog_service import load_catalog, load_chinese_core
    catalog = load_catalog()
    chinese_core = load_chinese_core()

    track_names = {r.get("期刊名称") for r in track_rows}
    all_rows = list(track_rows) if scope == "all" else []
    seen_names = set(track_names) if scope == "all" else set()

    for j in catalog.get("journals", []):
        name = j.get("name")
        if not name or name in seen_names:
            continue
        seen_names.add(name)
        level_code = j.get("level", "B")
        level_label = LEVEL_LABEL.get(level_code, j.get("category", "B类"))
        issn_parts = []
        if j.get("issn"):
            issn_parts.append(f"ISSN {j['issn']}")
        if j.get("cn"):
            issn_parts.append(f"CN {j['cn']}")

        all_rows.append({
            "期刊名称": name,
            "ISSN / CN": " | ".join(issn_parts),
            "收录索引": "广中医重要期刊目录（校办〔2020〕308号）",
            "广中医认定级别": level_label,
            "影响因子(JIF)": "官方直录",
            "中科院分区": "-",
            "主编信息": "-",
            "收稿方向": f"{j.get('category', '')}；主办单位：{j.get('publisher', '')}",
            "审稿周期": "同行评议",
            "出版方式与版面费": "传统订阅 / 常规版面费",
            "投稿方式": j.get("remark", ""),
            "核验日期与来源": "校办〔2020〕308号官方目录",
        })

    for cj in chinese_core:
        name = cj.get("name")
        if not name or name in seen_names:
            continue
        seen_names.add(name)
        calc_level = "A类2级" if cj.get("cscd_core") else ("A类3级" if cj.get("pkucore") else "A类4级")
        all_rows.append({
            "期刊名称": name,
            "ISSN / CN": f"ISSN {cj.get('issn', '')} | CN {cj.get('cn', '')}",
            "收录索引": "全国中文核心 / CSCD数据库收录",
            "广中医认定级别": calc_level,
            "影响因子(JIF)": "核心收录",
            "中科院分区": "-",
            "主编信息": cj.get("editor", {}).get("name", "-"),
            "收稿方向": f"学科分类：{cj.get('subject', '')}；{cj.get('scope', '')}",
            "审稿周期": cj.get("review_cycle", "同行评议"),
            "出版方式与版面费": cj.get("publication_mode", "传统订阅制"),
            "投稿方式": cj.get("submission_url", ""),
            "核验日期与来源": "中文核心/CSCD官方数据",
        })

    return all_rows
